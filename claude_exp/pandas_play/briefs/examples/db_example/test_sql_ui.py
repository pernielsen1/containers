"""Tests for sql_ui.py (see briefs/sql_ui.md).

Run:  python3 -m unittest test_sql_ui -v   (from db_example/)
"""
import json
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import sql_ui


class UiCase(unittest.TestCase):
    timeout = 600

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.scripts = self.dir / "sql"
        self.scripts.mkdir()
        self.cfg = self.dir / "config.json"
        self.cfg.write_text(json.dumps({"db_storage_dir": str(self.dir / "db")}))
        self.server = sql_ui.create_server(self.scripts, port=0, config=self.cfg,
                                           timeout=self.timeout)
        self.port = self.server.server_address[1]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.tmp.cleanup()

    def call(self, method, path, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=data, method=method,
                                     headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read() or b"null") if r.headers.get_content_type() == "application/json" else r.read()
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"null")

    def run_body(self, name="s.sql", text="PRAGMA print = 'hello';", **kw):
        body = {"name": name, "text": text, "env": "test", "db": "t"}
        body.update(kw)
        return body


class TestScripts(UiCase):
    def test_list_only_sql_sorted(self):
        (self.scripts / "b.sql").write_text("x")
        (self.scripts / "a.sql").write_text("x")
        (self.scripts / "notes.txt").write_text("x")
        (self.scripts / "sub").mkdir()
        (self.scripts / "sub" / "c.sql").write_text("x")
        status, data = self.call("GET", "/api/scripts")
        self.assertEqual(status, 200)
        self.assertEqual(data["scripts"], ["a.sql", "b.sql", "sub/c.sql"])
        self.assertEqual(data["dirs"], ["sub"])
        self.assertEqual(data["db"], "download")

    def test_load_and_missing(self):
        (self.scripts / "a.sql").write_text("select 1;", encoding="utf-8")
        status, data = self.call("GET", "/api/script?name=a.sql")
        self.assertEqual((status, data["text"]), (200, "select 1;"))
        status, _ = self.call("GET", "/api/script?name=nope.sql")
        self.assertEqual(status, 404)

    def test_save_creates_and_overwrites(self):
        self.assertEqual(self.call("POST", "/api/script", {"name": "n.sql", "text": "one"})[0], 200)
        self.assertEqual((self.scripts / "n.sql").read_text(encoding="utf-8"), "one")
        self.call("POST", "/api/script", {"name": "n.sql", "text": "two"})
        self.assertEqual((self.scripts / "n.sql").read_text(encoding="utf-8"), "two")

    def test_bad_names_rejected(self):
        for name in ["../x.sql", "sub/../../x.sql", "/etc/x.sql", "sub//x.sql", ".hidden/x.sql", "sub/.x.sql",
                     "x.txt", "", ".sql", "..\\x.sql", "sub\\x.sql"]:
            with self.subTest(name=name):
                self.assertEqual(self.call("POST", "/api/script", {"name": name, "text": "x"})[0], 400)
                self.assertEqual(self.call("GET", "/api/script?name=" + urllib.parse.quote(name))[0], 400)
        self.assertEqual(list(self.dir.glob("x.*")), [])

    def test_subfolder_save_load_run(self):
        (self.scripts / "more_sql").mkdir()
        self.assertEqual(self.call("POST", "/api/script", {"name": "more_sql/n.sql", "text": "hi"})[0], 200)
        self.assertEqual((self.scripts / "more_sql" / "n.sql").read_text(), "hi")
        status, data = self.call("GET", "/api/script?name=more_sql/n.sql")
        self.assertEqual((status, data["text"]), (200, "hi"))
        (self.scripts / "more_sql" / "inc.sql").write_text("PRAGMA print = 'sub include';")
        status, data = self.call("POST", "/api/run", self.run_body(
            name="more_sql/r.sql", text="PRAGMA include = 'inc.sql';"))
        self.assertEqual((status, data["exit"]), (200, 0))
        self.assertIn("sub include", data["output"])

    def test_empty_folder_listed_and_no_mkdir(self):
        (self.scripts / "empty").mkdir()
        _, data = self.call("GET", "/api/scripts")
        self.assertEqual((data["dirs"], data["scripts"]), (["empty"], []))
        self.assertEqual(self.call("POST", "/api/script", {"name": "newdir/x.sql", "text": "x"})[0], 400)
        self.assertFalse((self.scripts / "newdir").exists())

    def test_body_too_large(self):
        status, _ = self.call("POST", "/api/script", {"name": "big.sql", "text": "x" * (1024 * 1024 + 1)})
        self.assertEqual(status, 413)

    def test_binds_localhost_only(self):
        self.assertEqual(self.server.server_address[0], "127.0.0.1")


class TestRun(UiCase):
    def test_run_print(self):
        status, data = self.call("POST", "/api/run", self.run_body())
        self.assertEqual(status, 200)
        self.assertEqual(data["exit"], 0)
        self.assertIn("hello", data["output"])
        self.assertIn("seconds", data)

    def test_run_sql_error_is_200_nonzero(self):
        status, data = self.call("POST", "/api/run", self.run_body(text="SELEKT 1;"))
        self.assertEqual(status, 200)
        self.assertNotEqual(data["exit"], 0)
        self.assertIn("SELEKT", data["output"])

    def test_run_saves_posted_text_first(self):
        (self.scripts / "s.sql").write_text("PRAGMA print = 'old';")
        _, data = self.call("POST", "/api/run", self.run_body(text="PRAGMA print = 'new';"))
        self.assertIn("new", data["output"])
        self.assertNotIn("old", data["output"])
        self.assertIn("new", (self.scripts / "s.sql").read_text())

    def test_bad_env_db(self):
        self.assertEqual(self.call("POST", "/api/run", self.run_body(env="dev"))[0], 400)
        self.assertEqual(self.call("POST", "/api/run", self.run_body(env=""))[0], 400)
        body = self.run_body()
        del body["env"]
        self.assertEqual(self.call("POST", "/api/run", body)[0], 400)
        self.assertEqual(self.call("POST", "/api/run", self.run_body(db="a;b"))[0], 400)
        self.assertEqual(self.call("POST", "/api/run", self.run_body(db="../x"))[0], 400)

    def test_missing_db_defaults_to_download(self):
        body = self.run_body(text="CREATE TABLE t(a); INSERT INTO t VALUES (1);")
        del body["db"]
        status, data = self.call("POST", "/api/run", body)
        self.assertEqual((status, data["exit"]), (200, 0))
        self.assertTrue((self.dir / "db" / "download.db").exists())

    def test_relative_include(self):
        (self.scripts / "shared.sql").write_text("PRAGMA print = 'from shared';")
        _, data = self.call("POST", "/api/run",
                            self.run_body(text="PRAGMA include = 'shared.sql';"))
        self.assertEqual(data["exit"], 0)
        self.assertIn("from shared", data["output"])


class TestRunTable(UiCase):
    SELECT3 = "CREATE TABLE IF NOT EXISTS t(a, b); DELETE FROM t; INSERT INTO t VALUES (1,'x'),(2,'y'),(3,'z'); SELECT * FROM t"

    def test_table_returned(self):
        status, data = self.call("POST", "/api/run_table", self.run_body(text=self.SELECT3))
        self.assertEqual((status, data["exit"]), (200, 0))
        self.assertEqual(data["table"], {"columns": ["a", "b"], "rows": [["1", "x"], ["2", "y"], ["3", "z"]],
                                         "total": 3, "truncated": False})

    def test_trailing_semicolon_comment_and_no_semicolon_all_work(self):
        for tail in [";", "", "; -- done", " -- no semicolon yet"]:
            with self.subTest(tail=tail):
                _, data = self.call("POST", "/api/run_table",
                                    self.run_body(text="SELECT 1 AS n" + tail))
                self.assertEqual(data["exit"], 0, data["output"])
                self.assertEqual(data["table"]["rows"], [["1"]])

    def test_row_cap(self):
        text = "WITH RECURSIVE c(x) AS (SELECT 1 UNION ALL SELECT x+1 FROM c WHERE x < 1500) SELECT x FROM c;"
        _, data = self.call("POST", "/api/run_table", self.run_body(text=text))
        self.assertEqual((len(data["table"]["rows"]), data["table"]["total"], data["table"]["truncated"]),
                         (sql_ui.MAX_ROWS, 1500, True))

    def test_no_result_set_is_error_without_table(self):
        status, data = self.call("POST", "/api/run_table", self.run_body(text="PRAGMA print = 'hi';"))
        self.assertEqual(status, 200)
        self.assertNotEqual(data["exit"], 0)
        self.assertIn("no result set", data["output"])
        self.assertNotIn("table", data)

    def test_sql_error_no_table(self):
        _, data = self.call("POST", "/api/run_table", self.run_body(text="SELEKT 1;"))
        self.assertNotEqual(data["exit"], 0)
        self.assertNotIn("table", data)

    def test_saves_script_but_leaves_no_temp_files(self):
        (self.scripts / "sub").mkdir()
        (self.scripts / "sub" / "inc.sql").write_text("PRAGMA print = 'sub include';")
        text = "PRAGMA include = 'inc.sql';\nSELECT 7 AS n;"
        _, data = self.call("POST", "/api/run_table", self.run_body(name="sub/r.sql", text=text))
        self.assertEqual(data["exit"], 0, data["output"])          # relative include resolved
        self.assertEqual((self.scripts / "sub" / "r.sql").read_text(), text)   # no export pragma saved
        self.assertEqual(sorted(p.name for p in (self.scripts / "sub").iterdir()), ["inc.sql", "r.sql"])
        _, listing = self.call("GET", "/api/scripts")
        self.assertEqual(listing["scripts"], ["sub/inc.sql", "sub/r.sql"])

    def test_bad_env_and_name_rejected(self):
        self.assertEqual(self.call("POST", "/api/run_table", self.run_body(env="dev"))[0], 400)
        self.assertEqual(self.call("POST", "/api/run_table", self.run_body(name="../x.sql"))[0], 400)


class TestExportXlsx(UiCase):
    def load(self, body):
        import io
        from openpyxl import load_workbook
        status, data = self.call("POST", "/api/export_xlsx", body)
        self.assertEqual(status, 200)
        return list(load_workbook(io.BytesIO(data)).active.values)

    def test_header_rows_and_number_conversion(self):
        rows = self.load({"columns": ["n", "d", "id", "s"],
                          "rows": [["1", "1,5", "007", "x"], ["-2", "-0,25", "10", ""]]})
        self.assertEqual(rows[0], ("n", "d", "id", "s"))
        self.assertEqual(rows[1], (1, 1.5, "007", "x"))     # leading zero stays text
        self.assertEqual(rows[2][:3], (-2, -0.25, 10))

    def test_text_never_becomes_formula_or_crashes(self):
        # '=...' text used to be written as a formula (Excel: "file needs repair");
        # text containing a comma used to crash; control chars used to give a 500
        rows = self.load({"columns": ["=h"],
                          "rows": [["=foo bar"], ["=SUM(1;2)"], ["a,b"], ["x\x07y"], ["1e5"],
                                   ["12345678901234567890"], ["123456789012345"]]})
        self.assertEqual([r[0] for r in rows],
                         ["=h", "=foo bar", "=SUM(1;2)", "a,b", "xy", "1e5",
                          "12345678901234567890", 123456789012345])

    def test_no_formula_cells_in_xml(self):
        import io
        import zipfile
        status, data = self.call("POST", "/api/export_xlsx", {"columns": ["a"], "rows": [["=foo bar"]]})
        xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
        self.assertNotIn("<f>", xml)

    def test_core_properties_dates_are_valid_w3cdtf(self):
        # '2026-10-26T21:10:32+00:00Z' (offset AND Z) made Excel say "file needs repair"
        import io
        import re
        import zipfile
        _, data = self.call("POST", "/api/export_xlsx", {"columns": ["a"], "rows": [["1"]]})
        core = zipfile.ZipFile(io.BytesIO(data)).read("docProps/core.xml").decode()
        stamps = re.findall(r"<dcterms:(?:created|modified)[^>]*>([^<]*)<", core)
        self.assertEqual(len(stamps), 2)
        for st in stamps:
            self.assertRegex(st, r"^\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ$")

    def test_bad_payload_400(self):
        for body in [{}, {"columns": "a", "rows": []}, {"columns": ["a"], "rows": [[1]]},
                     {"columns": ["a"], "rows": ["x"]}]:
            with self.subTest(body=body):
                self.assertEqual(self.call("POST", "/api/export_xlsx", body)[0], 400)


class TestRunLimits(UiCase):
    timeout = 1

    def slow_script(self):
        # recursive CTE: a long-running query, no external sleep needed
        return ("WITH RECURSIVE c(x) AS (SELECT 1 UNION ALL SELECT x+1 FROM c WHERE x < 500000000) "
                "SELECT count(*) FROM c;")

    def test_timeout_kills(self):
        t0 = time.time()
        status, data = self.call("POST", "/api/run", self.run_body(text=self.slow_script()))
        self.assertEqual(status, 200)
        self.assertEqual(data["exit"], -1)
        self.assertIn("TIMEOUT", data["output"])
        self.assertLess(time.time() - t0, 20)

    def test_concurrent_run_409(self):
        first = {}
        t = threading.Thread(target=lambda: first.update(
            r=self.call("POST", "/api/run", self.run_body(text=self.slow_script()))))
        t.start()
        time.sleep(0.4)
        status, _ = self.call("POST", "/api/run", self.run_body())
        self.assertEqual(status, 409)
        t.join()
        self.assertEqual(first["r"][1]["exit"], -1)


class TestBrowser(UiCase):
    timeout = 4

    def setUp(self):
        super().setUp()
        try:
            from playwright.sync_api import sync_playwright
        except ImportError:
            self.skipTest("playwright not installed")
        self.pw = sync_playwright().start()
        try:
            self.browser = self.pw.chromium.launch()
        except Exception as e:
            self.pw.stop()
            self.skipTest(f"chromium not available: {e}")
        self.page = self.browser.new_page()
        (self.scripts / "a.sql").write_text("PRAGMA print = 'smoke';", encoding="utf-8")
        self.page.goto(f"http://127.0.0.1:{self.port}/")

    def tearDown(self):
        self.browser.close()
        self.pw.stop()
        super().tearDown()

    def test_smoke_edit_save_run(self):
        p = self.page
        p.click("text=a.sql")
        self.assertIn("smoke", p.input_value("#editor"))
        self.assertTrue(p.is_disabled("#run"))          # no env chosen yet
        p.select_option("#env", "test")
        p.fill("#db", "t")
        p.fill("#editor", "PRAGMA print = 'edited';")
        self.assertIn("*", p.inner_text(".tab.active"))   # unsaved marker
        p.click("#save")
        p.wait_for_function("!document.querySelector('.tab.active').innerText.includes('*')")
        self.assertIn("edited", (self.scripts / "a.sql").read_text())
        p.click("#run")
        p.wait_for_function("document.querySelector('#status').innerText.includes('exit 0')")
        self.assertIn("edited", p.inner_text("#output"))

    def test_new_script_from_empty_folder(self):
        p = self.page
        p.click("#newtab")
        p.fill("#editor", "PRAGMA print = 'fresh';")
        self.assertTrue(p.is_disabled("#run"))
        p.once("dialog", lambda d: d.accept("fresh"))
        p.click("#saveas")
        p.wait_for_function("document.querySelector('.tab.active').innerText.startsWith('fresh.sql')")
        self.assertIn("fresh", (self.scripts / "fresh.sql").read_text())

    def test_run_held_open_across_ticks(self):
        p = self.page
        (self.scripts / "slow.sql").write_text(TestRunLimits.slow_script(None), encoding="utf-8")
        p.reload()
        p.click("text=slow.sql")
        p.select_option("#env", "test")
        p.fill("#db", "t")
        p.click("#run")
        p.wait_for_timeout(1500)                        # several ticks while running
        self.assertTrue(p.is_disabled("#run"))
        p.fill("#editor", p.input_value("#editor") + " ")   # UI still responsive
        p.wait_for_function("!document.querySelector('#run').disabled")

    def test_tabs_keep_own_text_and_output(self):
        p = self.page
        (self.scripts / "b.sql").write_text("PRAGMA print = 'bee';", encoding="utf-8")
        p.reload()
        p.select_option("#env", "test")
        p.fill("#db", "t")
        p.click("#side >> text=a.sql")
        p.click("#side >> text=b.sql")
        self.assertEqual(p.locator(".tab").count(), 2)
        p.click("#side >> text=a.sql")                   # refocuses, no third tab
        self.assertEqual(p.locator(".tab").count(), 2)
        p.fill("#editor", "PRAGMA print = 'a-edit';")    # unsaved edit in a
        p.click(".tab >> text=b.sql")
        self.assertIn("bee", p.input_value("#editor"))
        p.click("#run")
        p.wait_for_function("document.querySelector('#status').innerText.includes('exit 0')")
        self.assertIn("bee", p.inner_text("#output"))
        p.click(".tab >> text=a.sql")
        self.assertIn("a-edit", p.input_value("#editor"))   # edit survived the switch
        self.assertIn("*", p.inner_text(".tab.active"))
        self.assertEqual(p.inner_text("#output"), "")        # b's output not shown here
        self.assertIn("smoke", (self.scripts / "a.sql").read_text())   # not saved to disk

    def test_close_tab(self):
        p = self.page
        p.click("#side >> text=a.sql")
        p.click(".tab.active .x")
        self.assertEqual(p.locator(".tab").count(), 0)
        self.assertTrue(p.is_disabled("#saveas"))
        p.click("#side >> text=a.sql")
        p.fill("#editor", "changed")
        p.once("dialog", lambda d: d.dismiss())          # keep it
        p.click(".tab.active .x")
        self.assertEqual(p.locator(".tab").count(), 1)
        p.once("dialog", lambda d: d.accept())           # discard
        p.click(".tab.active .x")
        self.assertEqual(p.locator(".tab").count(), 0)
        self.assertIn("smoke", (self.scripts / "a.sql").read_text())

    def test_run_and_table(self):
        p = self.page
        (self.scripts / "q.sql").write_text("SELECT 1 AS n, 'x' AS s;", encoding="utf-8")
        p.reload()
        p.click("#side >> text=q.sql")
        self.assertTrue(p.is_disabled("#runtable"))     # no env yet
        p.select_option("#env", "test")
        p.fill("#db", "t")
        self.assertTrue(p.is_hidden("#grid"))
        p.click("#runtable")
        p.wait_for_selector("#grid table")
        self.assertEqual(p.locator("#grid th").all_inner_texts(), ["n", "s"])
        self.assertEqual(p.locator("#grid td").all_inner_texts(), ["1", "x"])
        p.click("#side >> text=a.sql")                  # other tab: no grid
        self.assertTrue(p.is_hidden("#grid"))
        p.click(".tab >> text=q.sql")                   # grid is back
        self.assertEqual(p.locator("#grid td").count(), 2)
        p.click("#run")                                 # plain Run clears the grid
        p.wait_for_function("document.querySelector('#status').innerText.includes('exit 0')")
        self.assertTrue(p.is_hidden("#grid"))

    def run_table(self, text, js=None):
        """js (picker stub etc.) runs right after the page reload."""
        p = self.page
        (self.scripts / "q.sql").write_text(text, encoding="utf-8")
        p.reload()
        if js:
            p.evaluate("() => {" + js + "}")      # statement block: evaluate would call a function value
        p.click("#side >> text=q.sql")
        p.select_option("#env", "test")
        p.fill("#db", "t")
        p.click("#runtable")
        p.wait_for_selector("#grid table")

    def test_export_download_fallback(self):
        p = self.page
        self.run_table("SELECT 'a;b' AS s, 'say \"hi\"' AS q, 1.5 AS n;",
                       "delete window.showSaveFilePicker")      # as in Firefox
        with p.expect_download() as d:
            p.click("#exportcsv")
        self.assertEqual(d.value.suggested_filename, "q.csv")
        raw = Path(d.value.path()).read_bytes()
        self.assertEqual(raw.decode("utf-8-sig"), 's;q;n\n"a;b";"say ""hi""";1,5\n')
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))            # utf-8-sig like run_sql.py
        with p.expect_download() as d:
            p.click("#exportxlsx")
        self.assertEqual(d.value.suggested_filename, "q.xlsx")
        from openpyxl import load_workbook
        import io
        rows = list(load_workbook(io.BytesIO(Path(d.value.path()).read_bytes())).active.values)
        self.assertEqual(rows, [("s", "q", "n"), ("a;b", 'say "hi"', 1.5)])

    def test_export_with_file_picker(self):
        p = self.page
        self.run_table("SELECT 1 AS n;", """window.showSaveFilePicker = async opts => {
            window.__opts = opts;
            return {createWritable: async () => ({
                write: async b => { window.__saved = await b.text(); }, close: async () => {}})};
        }""")
        p.click("#exportcsv")
        p.wait_for_function("window.__saved !== undefined")
        self.assertEqual(p.evaluate("window.__opts.suggestedName"), "q.csv")
        self.assertEqual(p.evaluate("window.__saved").lstrip("\ufeff"), "n\n1\n")

    def test_export_picker_cancel_is_silent(self):
        p = self.page
        self.run_table("SELECT 1 AS n;", """window.showSaveFilePicker = async () => {
            throw new DOMException('cancelled', 'AbortError'); }""")
        p.on("dialog", lambda d: self.fail("unexpected alert: " + d.message))
        p.click("#exportcsv")
        p.wait_for_timeout(300)

    def test_untitled_tab_save_disabled_until_save_as(self):
        p = self.page
        p.click("#newtab")
        p.fill("#editor", "x")
        self.assertTrue(p.is_disabled("#save"))
        self.assertTrue(p.is_disabled("#run"))


if __name__ == "__main__":
    unittest.main()
